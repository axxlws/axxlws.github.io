"""Extract a figure placed directly before a caption in a Word report.

Used only to preserve the original report visual at its embedded resolution.
"""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document


REL_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"


def image_rids(paragraph):
    return [blip.get(REL_EMBED) for blip in paragraph._p.xpath(".//a:blip") if blip.get(REL_EMBED)]


def context(doc: Document, index: int, radius: int = 4):
    start = max(0, index - radius)
    end = min(len(doc.paragraphs), index + 2)
    for pos in range(start, end):
        text = " ".join(doc.paragraphs[pos].text.split())
        rids = image_rids(doc.paragraphs[pos])
        if text or rids:
            print(f"  [{pos:04d}] text={text[:150]!r} rids={rids}")


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("document", type=Path)
    parser.add_argument("caption")
    parser.add_argument("--output", type=Path)
    parser.add_argument("--list-only", action="store_true")
    parser.add_argument("--all-captions", action="store_true")
    parser.add_argument("--backtrack", type=int, default=7)
    args = parser.parse_args()

    doc = Document(args.document)
    if args.all_captions:
        for idx, paragraph in enumerate(doc.paragraphs):
            text = " ".join(paragraph.text.split())
            if "gambar" in text.casefold():
                print(f"[{idx:04d}] {text}")
        return
    caption_lower = args.caption.casefold()
    matches = [
        idx for idx, p in enumerate(doc.paragraphs)
        if caption_lower in " ".join(p.text.split()).casefold()
    ]
    if not matches:
        raise SystemExit(f"Caption not found: {args.caption!r}")

    print(f"Found {len(matches)} matching caption(s) in {args.document.name}")
    for index in matches:
        print(f"Caption paragraph {index}: {' '.join(doc.paragraphs[index].text.split())}")
        context(doc, index)

    if args.list_only:
        return
    if not args.output:
        raise SystemExit("--output is required unless --list-only is used")

    chosen_rid = None
    chosen_pos = None
    for index in matches:
        for pos in range(index - 1, max(-1, index - args.backtrack - 1), -1):
            rids = image_rids(doc.paragraphs[pos])
            if rids:
                chosen_rid = rids[-1]
                chosen_pos = pos
                break
        if chosen_rid:
            break
    if not chosen_rid:
        raise SystemExit("No embedded image found before the matching caption")

    part = doc.part.related_parts[chosen_rid]
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_bytes(part.blob)
    print(
        f"Extracted figure from paragraph {chosen_pos} as {args.output.name} "
        f"({part.content_type}, {len(part.blob):,} bytes)"
    )


if __name__ == "__main__":
    main()
